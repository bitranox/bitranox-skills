"""End-to-end tests for scripts/convert_with_ai.py.

The script runs as a real subprocess with offline stand-ins for markitdown and
openai on its import path (see conftest.py), so no API is ever called and the
tests run in CI. A "pptx" for the markitdown stand-in is a text file where each
line reading PICTURE is one picture. The one test against the REAL markitdown
builds a genuine deck with python-pptx and is skipped when either is missing.
"""

import struct
import zlib

import pytest

FAKE_KEY = {"OPENROUTER_API_KEY": "FAKE-KEY-NOT-REAL"}


def _ai(runner, tmp_path, fakes, args, *, mode="ok", encoding="utf-8"):
    env = dict(FAKE_KEY, FAKE_OPENAI_MODE=mode, FAKE_OPENAI_LOG=str(tmp_path / "calls.log"))
    return runner("convert_with_ai", args, cwd=tmp_path, pythonpath=fakes, env=env, encoding=encoding)


def _calls(tmp_path):
    log = tmp_path / "calls.log"
    return len(log.read_text(encoding="utf-8").splitlines()) if log.exists() else 0


@pytest.fixture
def fakes(fake_markitdown_dir, fake_openai_dir):
    return [fake_markitdown_dir, fake_openai_dir]


def test_list_prompts_needs_no_positionals(script_runner, tmp_path, fakes):
    run = _ai(script_runner, tmp_path, fakes, ["--list-prompts"])

    assert run.returncode == 0, run.output
    assert "Available prompt types" in run.stdout


def test_missing_positionals_are_still_a_usage_error(script_runner, tmp_path, fakes):
    run = _ai(script_runner, tmp_path, fakes, [])

    assert run.returncode == 2, run.output
    assert "input" in run.stderr


def test_help_works_without_openai_or_markitdown(script_runner, block_import, tmp_path):
    run = script_runner(
        "convert_with_ai", ["--help"], cwd=tmp_path, prelude=block_import("markitdown", "openai"),
    )

    assert run.returncode == 0, run.output
    assert "usage" in run.stdout.lower()


def test_pep723_block_declares_both_dependencies(scripts_dir):
    text = (scripts_dir / "convert_with_ai.py").read_text(encoding="utf-8")
    block = text.split("# /// script", 1)[1].split("# ///", 1)[0]
    assert "markitdown[all]" in block
    assert "openai" in block


def _deck(path, pictures):
    path.write_text("SLIDE TEXT\n" + "PICTURE\n" * pictures, encoding="utf-8")


def test_pptx_whose_captions_all_fail_fails_the_run(script_runner, tmp_path, fakes):
    _deck(tmp_path / "deck.pptx", 1)

    run = _ai(script_runner, tmp_path, fakes, ["deck.pptx", "deck.md"], mode="fail")

    assert run.returncode == 1, run.output
    assert _calls(tmp_path) == 1
    assert "1 of 1 image description" in run.stderr
    assert "[OK] Successfully converted" not in run.output


def test_pptx_with_working_captions_succeeds(script_runner, tmp_path, fakes):
    _deck(tmp_path / "deck.pptx", 2)

    run = _ai(script_runner, tmp_path, fakes, ["deck.pptx", "deck.md"])

    assert run.returncode == 0, run.output
    assert (tmp_path / "deck.md").read_text(encoding="utf-8").count("FAKE-CAPTION") == 2


def test_png_whose_caption_fails_fails_the_run(script_runner, tmp_path, fakes):
    (tmp_path / "pic.png").write_text("PNG", encoding="utf-8")

    run = _ai(script_runner, tmp_path, fakes, ["pic.png", "pic.md"], mode="fail")

    assert run.returncode == 1, run.output
    assert not (tmp_path / "pic.md").exists()


def test_cp1252_console_with_non_ascii_input_name(script_runner, tmp_path, fakes):
    name = "Grafik_" + chr(0xFC) + "_" + chr(0x65E5) + chr(0x672C) + ".png"
    (tmp_path / name).write_text("PNG", encoding="utf-8")

    run = _ai(script_runner, tmp_path, fakes, [name, "out.md"], encoding="cp1252")

    assert run.returncode == 0, run.output
    assert "FAKE-CAPTION" in (tmp_path / "out.md").read_text(encoding="utf-8")


def test_cp1252_console_with_non_ascii_output_name(script_runner, tmp_path, fakes):
    (tmp_path / "pic.png").write_text("PNG", encoding="utf-8")
    out_name = chr(0x65E5) + chr(0x672C) + ".md"

    run = _ai(script_runner, tmp_path, fakes, ["pic.png", out_name], encoding="cp1252")

    assert run.returncode == 0, run.output
    assert (tmp_path / out_name).is_file()


@pytest.mark.parametrize("name", ["paper.pdf", "notes.DOCX", "sheet.xlsx"])
def test_input_without_an_ai_description_path_is_refused(script_runner, tmp_path, fakes, name):
    (tmp_path / name).write_text("CONTENT", encoding="utf-8")

    run = _ai(script_runner, tmp_path, fakes, [name, "out.md"])

    assert run.returncode == 1, run.output
    assert "no AI image descriptions" in run.stderr
    assert not (tmp_path / "out.md").exists()
    assert _calls(tmp_path) == 0


@pytest.mark.parametrize("name", ["pic.PNG", "pic.jpg", "pic.jpeg"])
def test_image_suffixes_are_accepted_in_any_case(script_runner, tmp_path, fakes, name):
    (tmp_path / name).write_text("IMG", encoding="utf-8")

    run = _ai(script_runner, tmp_path, fakes, [name, "out.md"])

    assert run.returncode == 0, run.output
    assert _calls(tmp_path) == 1


def test_missing_api_key_exits_1(script_runner, tmp_path, fakes):
    (tmp_path / "pic.png").write_text("PNG", encoding="utf-8")

    run = script_runner("convert_with_ai", ["pic.png", "out.md"], cwd=tmp_path, pythonpath=fakes)

    assert run.returncode == 1, run.output
    assert "OPENROUTER_API_KEY" in run.stderr
    # The one remedy it names must be one that works: there is no key flag to point at.
    assert "--api-key" not in run.output


def test_a_missing_input_file_is_reported_on_stderr(script_runner, tmp_path, fakes):
    """The arm: the error went to STDOUT, where a caller capturing the progress lines (or piping
    them on) reads it as output and a caller watching stderr sees nothing."""
    run = _ai(script_runner, tmp_path, fakes, ["absent.png", "out.md"])

    assert run.returncode == 1, run.output
    assert "does not exist" in run.stderr
    assert "does not exist" not in run.stdout
    assert not (tmp_path / "out.md").exists()
    assert _calls(tmp_path) == 0


def test_a_directory_as_input_is_refused_as_not_a_file(script_runner, tmp_path, fakes):
    """A directory named like an image passed the existence check and reached markitdown, which
    failed with an error about the converter rather than about the input."""
    (tmp_path / "shots.png").mkdir()

    run = _ai(script_runner, tmp_path, fakes, ["shots.png", "out.md"])

    assert run.returncode == 1, run.output
    assert "is not a file" in run.stderr
    assert "is not a file" not in run.stdout
    assert _calls(tmp_path) == 0


def test_an_existing_input_file_passes_the_input_check(script_runner, tmp_path, fakes):
    """The control: the same call with the file present converts."""
    (tmp_path / "absent.png").write_text("PNG", encoding="utf-8")

    run = _ai(script_runner, tmp_path, fakes, ["absent.png", "out.md"])

    assert run.returncode == 0, run.output
    assert "does not exist" not in run.output


SECRET ="sk-or-v1-ARGV-SECRET-MUST-NOT-BE-USED"


@pytest.mark.parametrize(
    "key_args",
    [["--api-key", SECRET], ["-k", SECRET], [f"--api-key={SECRET}"], ["--api", SECRET]],
    ids=["long", "short", "equals", "abbreviated"],
)
@pytest.mark.parametrize("env_key", [False, True], ids=["no-env-key", "env-key-set"])
def test_an_api_key_on_the_command_line_is_refused(
    script_runner, tmp_path, fakes, key_args, env_key
):
    """A key in argv sits in the process list, shell history and CI logs for the whole run.

    Refused loudly rather than ignored, and refused even when the environment ALSO holds a key,
    so a caller learns the flag is gone instead of believing the argv key was the one used.
    """
    (tmp_path / "pic.png").write_text("PNG", encoding="utf-8")
    env = dict(FAKE_KEY) if env_key else {}
    env["FAKE_OPENAI_LOG"] = str(tmp_path / "calls.log")

    run = script_runner(
        "convert_with_ai", ["pic.png", "out.md", *key_args], cwd=tmp_path, pythonpath=fakes, env=env
    )

    assert run.returncode == 2, run.output
    assert "OPENROUTER_API_KEY" in run.stderr
    assert SECRET not in run.output
    assert _calls(tmp_path) == 0
    assert not (tmp_path / "out.md").exists()


def test_the_environment_key_alone_still_converts(script_runner, tmp_path, fakes):
    """The control for the refusal: the same run with the key where it belongs succeeds."""
    (tmp_path / "pic.png").write_text("PNG", encoding="utf-8")

    run = _ai(script_runner, tmp_path, fakes, ["pic.png", "out.md"])

    assert run.returncode == 0, run.output
    assert _calls(tmp_path) == 1


def test_help_offers_no_key_flag(script_runner, tmp_path):
    run = script_runner("convert_with_ai", ["--help"], cwd=tmp_path)

    assert run.returncode == 0, run.output
    assert "--api-key" not in run.stdout
    assert "-k," not in run.stdout
    assert "OPENROUTER_API_KEY" in run.stdout


def _png():
    """A valid 1x1 RGB PNG."""
    def chunk(tag, data):
        body = tag + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body))
    header = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    pixels = zlib.compress(b"\x00\xff\x00\x00")
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", header) + chunk(b"IDAT", pixels) + chunk(b"IEND", b"")


@pytest.mark.parametrize("mode,expected_rc", [("fail", 1), ("ok", 0)])
def test_real_markitdown_pptx_caption_failure_is_seen(script_runner, fake_openai_dir, tmp_path, mode, expected_rc):
    pytest.importorskip("markitdown")
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    (tmp_path / "pic.png").write_bytes(_png())
    deck = pptx.Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_picture(str(tmp_path / "pic.png"), Inches(1), Inches(1))
    deck.save(str(tmp_path / "deck.pptx"))

    run = _ai(script_runner, tmp_path, [fake_openai_dir], ["deck.pptx", "deck.md"], mode=mode)

    assert run.returncode == expected_rc, run.output
    assert _calls(tmp_path) == 1
